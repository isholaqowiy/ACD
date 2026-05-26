import os
import uuid
import json
import logging
from datetime import date
from aiogram import Router, F, Bot
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.fsm.context import FSMContext
from sqlalchemy.future import select
from database import AsyncSessionLocal, User, ScanHistory, DailyUsage
from keyboards import cancel_keyboard, export_keyboard, main_keyboard, back_keyboard
from services.ai_detector import analyze_text
from states import ScanStates

logger = logging.getLogger(__name__)
router = Router()

FREE_DAILY_LIMIT = 10

async def check_usage(user_id: int, is_premium: bool) -> bool:
    """Returns True if user can scan, False if limit reached."""
    if is_premium:
        return True
    async with AsyncSessionLocal() as session:
        today = date.today()
        result = await session.execute(
            select(DailyUsage).where(
                DailyUsage.telegram_id == user_id,
                DailyUsage.usage_date == today
            )
        )
        usage = result.scalar_one_or_none()
        if usage and usage.count >= FREE_DAILY_LIMIT:
            return False
    return True

async def increment_usage(user_id: int):
    async with AsyncSessionLocal() as session:
        today = date.today()
        result = await session.execute(
            select(DailyUsage).where(
                DailyUsage.telegram_id == user_id,
                DailyUsage.usage_date == today
            )
        )
        usage = result.scalar_one_or_none()
        if not usage:
            usage = DailyUsage(telegram_id=user_id, usage_date=today, count=0)
            session.add(usage)
        usage.count += 1
        await session.commit()

async def get_user(user_id: int) -> User | None:
    async with AsyncSessionLocal() as session:
        result = await session.execute(
            select(User).where(User.telegram_id == user_id)
        )
        return result.scalar_one_or_none()

async def run_scan(message: Message, state: FSMContext, text: str, file_name: str = "Text Input"):
    user_id = message.from_user.id
    user = await get_user(user_id)
    is_premium = user.is_premium if user else False

    if not await check_usage(user_id, is_premium):
        await state.clear()
        return await message.answer(
            "❌ *Daily Limit Reached (10/10)*\n\n"
            "Upgrade to Premium for unlimited scans.\n"
            "Use /subscription to learn more.",
            parse_mode="Markdown",
            reply_markup=main_keyboard()
        )

    loading = await message.answer("🔄 *Analyzing content...* Please wait.", parse_mode="Markdown")

    try:
        report = await analyze_text(text)
        ai_prob = int(report.get("ai_probability", 0))
        human_prob = 100 - ai_prob
        verdict = report.get("verdict", "Unknown")
        confidence = report.get("confidence", "Low")
        analysis = report.get("analysis", [])
        recommendation = report.get("recommendation", "")

        async with AsyncSessionLocal() as session:
            record = ScanHistory(
                telegram_id=user_id,
                file_name=file_name,
                ai_percent=ai_prob,
                human_percent=human_prob,
                verdict=verdict,
                details=json.dumps(report)
            )
            session.add(record)
            await session.commit()
            record_id = record.id

        await increment_usage(user_id)

        bullets = "\n".join([f"• _{b}_" for b in analysis])
        result_text = (
            f"📊 *AI Detection Result*\n"
            f"━━━━━━━━━━━━━━━━━━━━━\n"
            f"🤖 *AI:* `{ai_prob}%` | 🧑 *Human:* `{human_prob}%`\n"
            f"⚖️ *Verdict:* *{verdict}*\n"
            f"🎯 *Confidence:* `{confidence}`\n\n"
            f"💡 *Key Findings:*\n{bullets}\n\n"
            f"📋 *Recommendation:*\n_{recommendation}_"
        )

        await loading.delete()
        await message.answer(result_text, parse_mode="Markdown", reply_markup=export_keyboard(record_id))
        await state.clear()

    except Exception as e:
        logger.error(f"Scan error: {e}")
        await loading.delete()
        await state.clear()
        await message.answer(
            "❌ Analysis failed. Please try again.",
            reply_markup=main_keyboard()
        )

@router.callback_query(F.data == "detect")
async def detect_callback(callback: CallbackQuery, state: FSMContext):
    await state.set_state(ScanStates.waiting_for_content)
    try:
        await callback.message.edit_text(
            "📝 *Paste your text below.*\n\n"
            "Minimum 100 characters for accurate results.",
            parse_mode="Markdown",
            reply_markup=cancel_keyboard()
        )
    except Exception:
        pass
    await callback.answer()

@router.callback_query(F.data == "upload")
async def upload_callback(callback: CallbackQuery, state: FSMContext):
    user = await get_user(callback.from_user.id)
    if not user or not user.is_premium:
        try:
            await callback.message.edit_text(
                "💎 *Premium Required*\n\n"
                "File upload (PDF/DOCX) is a Premium feature.\n"
                "Use /subscription to upgrade.",
                parse_mode="Markdown",
                reply_markup=back_keyboard()
            )
        except Exception:
            pass
        return await callback.answer()

    await state.set_state(ScanStates.waiting_for_file)
    try:
        await callback.message.edit_text(
            "📂 *Upload your file.*\n\n"
            "Supported formats: `.pdf`, `.docx`, `.txt`",
            parse_mode="Markdown",
            reply_markup=cancel_keyboard()
        )
    except Exception:
        pass
    await callback.answer()

@router.message(ScanStates.waiting_for_content)
async def handle_text_input(message: Message, state: FSMContext):
    text = message.text or ""
    if len(text.strip()) < 100:
        return await message.answer(
            "⚠️ Text too short. Please provide at least 100 characters."
        )
    await run_scan(message, state, text)

@router.message(ScanStates.waiting_for_file)
async def handle_file_input(message: Message, state: FSMContext, bot: Bot):
    if not message.document:
        return await message.answer("⚠️ Please upload a valid file.")

    file_name = message.document.file_name or "upload"
    ext = os.path.splitext(file_name)[1].lower()

    if ext not in [".pdf", ".docx", ".txt"]:
        return await message.answer("⚠️ Only PDF, DOCX, and TXT files are supported.")

    loading = await message.answer("📥 *Downloading file...*", parse_mode="Markdown")

    try:
        file_info = await bot.get_file(message.document.file_id)
        temp_path = f"/tmp/{uuid.uuid4()}{ext}"
        await bot.download_file(file_info.file_path, destination=temp_path)

        if ext == ".txt":
            with open(temp_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext == ".pdf":
            import PyPDF2
            with open(temp_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(temp_path)
            text = "\n".join(p.text for p in doc.paragraphs)

        os.remove(temp_path)
        await loading.delete()

        if len(text.strip()) < 100:
            await state.clear()
            return await message.answer(
                "⚠️ File content too short for analysis.",
                reply_markup=main_keyboard()
            )

        await run_scan(message, state, text, file_name)

    except Exception as e:
        logger.error(f"File processing error: {e}")
        await loading.delete()
        await state.clear()
        await message.answer("❌ Failed to process file. Please try again.", reply_markup=main_keyboard())

@router.callback_query(F.data.startswith("export_"))
async def export_handler(callback: CallbackQuery):
    parts = callback.data.split("_")
    fmt = parts[1]
    scan_id = int(parts[2])

    async with AsyncSessionLocal() as session:
        record = await session.get(ScanHistory, scan_id)

    if not record:
        return await callback.answer("❌ Record not found.", show_alert=True)

    details = json.loads(record.details)
    out_path = f"/tmp/report_{scan_id}.{fmt}"

    if fmt == "txt":
        with open(out_path, "w") as f:
            f.write(f"AI Detection Report\n")
            f.write(f"==================\n")
            f.write(f"File: {record.file_name}\n")
            f.write(f"Date: {record.scanned_at.strftime('%Y-%m-%d %H:%M')}\n")
            f.write(f"Verdict: {record.verdict}\n")
            f.write(f"AI: {record.ai_percent}% | Human: {record.human_percent}%\n\n")
            f.write(f"Analysis:\n")
            for b in details.get("analysis", []):
                f.write(f"- {b}\n")
            f.write(f"\nRecommendation:\n{details.get('recommendation', '')}\n")
    else:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, "AI Detection Report", ln=True)
        pdf.cell(0, 10, f"File: {record.file_name}", ln=True)
        pdf.cell(0, 10, f"Date: {record.scanned_at.strftime('%Y-%m-%d %H:%M')}", ln=True)
        pdf.cell(0, 10, f"Verdict: {record.verdict}", ln=True)
        pdf.cell(0, 10, f"AI: {record.ai_percent}% | Human: {record.human_percent}%", ln=True)
        pdf.cell(0, 10, "Analysis:", ln=True)
        for b in details.get("analysis", []):
            pdf.multi_cell(0, 8, f"- {b}")
        pdf.cell(0, 10, "Recommendation:", ln=True)
        pdf.multi_cell(0, 8, details.get("recommendation", ""))
        pdf.output(out_path)

    await callback.message.answer_document(
        FSInputFile(out_path),
        caption="📄 Your detection report."
    )
    await callback.answer()
    if os.path.exists(out_path):
        os.remove(out_path)

@router.callback_query(F.data == "history")
async def history_handler(callback: CallbackQuery):
    user_id = callback.from_user.id
    async with AsyncSessionLocal() as session:
        result = await session.execute(
            select(ScanHistory)
            .where(ScanHistory.telegram_id == user_id)
            .order_by(ScanHistory.scanned_at.desc())
            .limit(5)
        )
        records = result.scalars().all()

    if not records:
        try:
            await callback.message.edit_text(
                "📜 *No scan history yet.*\n\nRun your first scan to see results here.",
                parse_mode="Markdown",
                reply_markup=main_keyboard()
            )
        except Exception:
            pass
        return await callback.answer()

    text = "📜 *Recent Scans (Last 5):*\n\n"
    for r in records:
        text += (
            f"• `{r.scanned_at.strftime('%Y-%m-%d')}` — *{r.verdict}*\n"
            f"  AI: `{r.ai_percent}%` | File: _{r.file_name[:30]}_\n\n"
        )
    try:
        await callback.message.edit_text(text, parse_mode="Markdown", reply_markup=main_keyboard())
    except Exception:
        pass
    await callback.answer()
